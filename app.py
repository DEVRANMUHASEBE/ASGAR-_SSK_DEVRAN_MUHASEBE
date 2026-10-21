import io, os, json, calendar, hashlib
from datetime import date
from pathlib import Path
import streamlit as st
import pandas as pd
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, Image
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Resmi Gazete yapi yaklasik birim maliyetleri (TL/m2).
# 2024 RG 20.02.2024/32466; 2025 RG 31.01.2025/32799;
# 2026 RG 03.02.2026/33157.
OFFICIAL_COSTS = {
    2024: {'I-A':1500,'I-B':2100,'II-A':3500,'II-B':5250,'II-C':7750,
           'III-A':12250,'III-B':14400,'III-C':None,'IV-A':15300,
           'IV-B':17400,'IV-C':18700,'V-A':21300,'V-B':22250,
           'V-C':24300,'V-D':26800},
    2025: {'I-A':2100,'I-B':3050,'I-C':3300,'I-D':3900,
           'II-A':6600,'II-B':10200,'II-C':12400,'III-A':17100,
           'III-B':18200,'III-C':19150,'IV-A':21500,'IV-B':27500,
           'IV-C':32600,'V-A':34400,'V-B':35600,'V-C':39500,
           'V-D':43400,'V-E':86250},
    2026: {'I-A':2600,'I-B':3900,'I-C':4200,'I-D':4800,
           'II-A':8100,'II-B':12500,'II-C':15100,'III-A':19800,
           'III-B':21050,'III-C':23400,'IV-A':26450,'IV-B':33900,
           'IV-C':40500,'V-A':42350,'V-B':43850,'V-C':48750,
           'V-D':53500,'V-E':103500},
}

def tariff_years(entry):
    """Bitiş yılı hariç: 1-2 takvim yılında başlangıç yılı; 3+ yılda bitişten önceki yılların ortalaması."""
    first, last = entry['start'].year, entry['end'].year
    if last < first:
        return []
    if last - first <= 1:
        return [first]
    return list(range(first, last))

def tariff_cost(year, klass):
    return OFFICIAL_COSTS.get(year, {}).get(klass)

st.set_page_config(page_title="DEVRAN | Asgari İşçilik",page_icon="🏗️",layout="wide")
st.markdown('''<style>.stApp{background:#f7f9fc} .block-container{padding-top:1.4rem;max-width:1320px}h1,h2,h3{color:#172d4f}.stButton button[kind=primary]{background:#16345e;border-color:#16345e}</style>''',unsafe_allow_html=True)
LOGO = Path(__file__).with_name('logo.png')
a,b=st.columns([1,7])
with a:
    if LOGO.exists(): st.image(str(LOGO),width=110)
    else: st.markdown('## 🏛️')
with b:
    st.title('DEVRAN MÂLİ MÜŞAVİRLİK')
    st.caption('İNŞAAT ASGARİ İŞÇİLİK HESAPLAMA • v1.3')
st.divider()

def tr(n, unit=' TL'):
    return f'{n:,.2f}'.replace(',','X').replace('.',',').replace('X','.')+unit

def money(v):
    try:
        if isinstance(v,str): return float(v.replace('.','').replace(',','.'))
        return float(v or 0)
    except (ValueError, TypeError): return 0.0

def build_xlsx(project, details, monthly, total, required, base, missing, estimated, rate, prime):
    """Tek sayfada proje özeti ve altında kullanıcı tarafından ömür boyu doldurulan aylık SGK kayıtları."""
    wb=Workbook()
    ws=wb.active; ws.title='İNŞAAT SGK TAKİP'
    ws.sheet_view.showGridLines=False
    for col,width in {'A':14,'B':18,'C':18,'D':19,'E':19,'F':22,'G':22,'H':19}.items():
        ws.column_dimensions[col].width=width
    navy='16345E'; pale='E8EEF6'; light='F4F7FB'; blue='DDEBFB'; white='FFFFFF'
    ws.merge_cells('A1:H2'); ws['A1']='DEVRAN MÂLİ MÜŞAVİRLİK  |  İNŞAAT ASGARİ İŞÇİLİK / SGK TAKİP'
    ws['A1'].fill=PatternFill('solid',fgColor=navy);ws['A1'].font=Font(name='Calibri',size=15,bold=True,color=white)
    ws['A1'].alignment=Alignment(vertical='center',horizontal='center');ws.row_dimensions[1].height=25;ws.row_dimensions[2].height=18
    def field(row,key,val,col='A',end='D'):
        labels={'A':('A','B'),'E':('E','F')}; lc,vc=labels[col]
        ws[f'{lc}{row}']=key; ws[f'{lc}{row}'].font=Font(bold=True,color=navy)
        from openpyxl.utils import column_index_from_string
        target_right='D' if col=='A' else 'H'
        ws.merge_cells(f'{vc}{row}:{target_right}{row}');ws[f'{vc}{row}']=val
        ws[f'{vc}{row}'].alignment=Alignment(vertical='center')
        ws.row_dimensions[row].height=24
    field(4,'Proje / Mükellef',project['ad']);field(4,'Ada / Parsel',project['ada']+' / '+project['parsel'],'E')
    field(5,'İnşaat dönemi',project['dates']);field(5,'Yapı / Blok sayısı',len(details),'E')
    field(6,'Asgari işçilik oranı',rate/100); ws['B6'].number_format='0.00%'
    field(6,'Prim oranı',prime/100,'E');ws['F6'].number_format='0.00%'
    field(7,'Araştırma indirimi',project['discount']/100);ws['B7'].number_format='0.00%'
    ws.merge_cells('A9:H9');ws['A9']='HESAP ÖZETİ';ws['A9'].fill=PatternFill('solid',fgColor=navy);ws['A9'].font=Font(bold=True,color=white)
    values=[('Toplam inşaat alanı (m²)',sum(d['area'] for d in details)),('Toplam inşaat maliyeti (TL)',total),('Gerekli asgari işçilik matrahı (TL)',required),('Bildirilen PEK toplamı (TL)',None),('Eksik işçilik matrahı (TL)',None),('Tahmini ek prim (TL)',None),('Ödenen SGK primi toplamı (TL)',None)]
    for idx,(label,val) in enumerate(values,10):
        ws.merge_cells(start_row=idx,start_column=1,end_row=idx,end_column=4)
        ws.cell(idx,1,label);ws.cell(idx,1).font=Font(bold=True,color=navy)
        ws.merge_cells(start_row=idx,start_column=5,end_row=idx,end_column=8)
        ws.cell(idx,5,val);ws.cell(idx,5).number_format='#,##0.00';ws.row_dimensions[idx].height=23
        if idx%2==0:
            for cells in ws.iter_rows(min_row=idx,max_row=idx,min_col=1,max_col=8):
                for c in cells:c.fill=PatternFill('solid',fgColor=light)
    ws.merge_cells('A18:H18');ws['A18']='YAPI / BLOK DETAYI';ws['A18'].fill=PatternFill('solid',fgColor=navy);ws['A18'].font=Font(bold=True,color=white)
    hdr=19
    for c,v in enumerate(['Yapı / Blok','Sınıf','Alan m²','Tarife yılları','Birim TL/m²','Yapı maliyeti TL'],1):
        ws.cell(hdr,c,v);ws.cell(hdr,c).font=Font(bold=True,color=navy);ws.cell(hdr,c).fill=PatternFill('solid',fgColor=pale)
    for i,d in enumerate(details,20):
        for c,v in enumerate([d['name'],d['klass'],d['area'],', '.join(map(str,d['years'])),d['cost'],d['total']],1):ws.cell(i,c,v)
        for c in (3,5,6):ws.cell(i,c).number_format='#,##0.00'
    begin=22+len(details)
    ws.merge_cells(start_row=begin,start_column=1,end_row=begin,end_column=8)
    ws.cell(begin,1,'AYLIK SGK TAKİP - MAVİ ALANLARI HER AY DOLDURUN');ws.cell(begin,1).fill=PatternFill('solid',fgColor=navy);ws.cell(begin,1).font=Font(bold=True,color=white)
    labels=['Yıl','Ay','Prime Esas Kazanç (TL)','Ödenen SGK Primi (TL)','Açıklama','Ödeme tarihi','Durum','Not']
    for c,label in enumerate(labels,1):
        z=ws.cell(begin+1,c,label);z.font=Font(bold=True,color=navy);z.fill=PatternFill('solid',fgColor=pale);z.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[begin+1].height=36
    first=begin+2
    for i,r in enumerate(monthly):
        row=first+i
        for c,v in enumerate([int(r[0]),str(r[1]),float(r[2]),float(r[3]),'','','',''],1):
            cell=ws.cell(row,c,v);cell.fill=PatternFill('solid',fgColor=blue if c>=3 else light)
            if c in (3,4):cell.number_format='#,##0.00'
        ws.row_dimensions[row].height=22
    last=first+len(monthly)-1
    if monthly:
        ws['E13']=f'=SUM(C{first}:C{last})'
        ws['E14']='=MAX(0,E12-E13)'
        ws['E15']='=E14*F6'
        ws['E16']=f'=SUM(D{first}:D{last})'
    ws.merge_cells(start_row=last+3,start_column=1,end_row=last+3,end_column=8)
    ws.cell(last+3,1,'BİLGİ: Eksik işçilik karşılaştırmasında ödenen SGK primi değil, bildirilen prime esas kazanç kullanılır.')
    ws.cell(last+3,1).font=Font(italic=True,color=navy,size=10)
    ws.freeze_panes=f'C{first}';ws.print_options.horizontalCentered=True
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_area=f'A1:H{last+3}'
    buf=io.BytesIO();wb.save(buf);return buf.getvalue()

def read_xlsx(data):
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    sh=wb['İNŞAAT SGK TAKİP'] if 'İNŞAAT SGK TAKİP' in wb.sheetnames else wb['Aylik_SGK']
    result=[];in_table=False
    months_set={'Ocak','Şubat','Mart','Nisan','Mayıs','Haziran','Temmuz','Ağustos','Eylül','Ekim','Kasım','Aralık'}
    for r in sh.iter_rows(values_only=True):
        if len(r)>1 and str(r[0]).strip()=='Yıl' and str(r[1]).strip()=='Ay': in_table=True;continue
        if not in_table:continue
        if len(r)<4 or not isinstance(r[0],(int,float)) or r[1] not in months_set:continue
        result.append([int(r[0]),str(r[1]),money(r[2]),money(r[3])])
    wb.close();return result

def pdf_report(project,details,monthly,total,required,base,missing,estimated,rate,prime):
    buf=io.BytesIO(); doc=SimpleDocTemplate(buf,pagesize=A4,rightMargin=34,leftMargin=34,topMargin=32,bottomMargin=32)
    fonts=[str(Path(__file__).with_name('DejaVuSans.ttf')),'/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf','/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf']
    font=next((x for x in fonts if os.path.exists(x)),None)
    fname='Helvetica'
    if font:
        pdfmetrics.registerFont(TTFont('Devran',font));fname='Devran'
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='DevTitle',fontName=fname,fontSize=15,leading=20,alignment=TA_CENTER,spaceAfter=14))
    styles.add(ParagraphStyle(name='DevTxt',fontName=fname,fontSize=9,leading=14))
    story=[]
    if LOGO.exists():
        try:
            logo=Image(str(LOGO),width=68,height=68,kind='proportional');logo.hAlign='CENTER';story.append(logo);story.append(Spacer(1,5))
        except Exception:pass
    story += [Paragraph('DEVRAN MÂLİ MÜŞAVİRLİK',styles['DevTitle']),Paragraph('İNŞAAT ASGARİ İŞÇİLİK - ÖZET HESAP RAPORU',styles['DevTxt']),Spacer(1,12)]
    info=[['Proje / Mükellef',project['ad']],['Ada / Parsel',project['ada']+' / '+project['parsel']],['İnşaat Dönemi',project['dates']],['Yapı / Blok Sayısı',str(len(details))]]
    def tbl(rows,widths):
        t=Table([[Paragraph(str(v),styles['DevTxt']) for v in row] for row in rows],colWidths=widths,repeatRows=1)
        t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.35,colors.lightgrey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#E8EDF5')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]))
        return t
    story += [tbl(info,[145,365]),Spacer(1,12)]
    summary=[['Kalem','Tutar'],['Toplam inşaat maliyeti',tr(total)],['Gerekli asgari işçilik matrahı',tr(required)],['Bildirilen prime esas kazanç',tr(base)],['Eksik işçilik matrahı',tr(missing)],['Tahmini ek prim (%'+tr(prime,'')+')',tr(estimated)],['Fiilen ödenen SGK primi (bilgi)',tr(sum(x[3] for x in monthly))]]
    story += [tbl(summary,[320,190]),Spacer(1,12)]
    detailrows=[['Blok / Yapı','Alan (m²)','Birim maliyet','Yapı maliyeti']]
    for r in details: detailrows.append([r['name'],tr(r['area'],''),tr(r['cost'],''),tr(r['total'],'')])
    story += [tbl(detailrows,[146,99,125,140]),Spacer(1,12)]
    explanation=f"Asgari işçilik oranı %{rate:g}, araştırma indirimi %{project['discount']:g}; tahmini prim oranı %{prime:g}. Birim maliyet ve tarihler kontrol edilmelidir. Prim tahmini gecikme zammı, ceza ve teşvikleri içermez."
    story.append(Paragraph(explanation,styles['DevTxt']))
    doc.build(story);return buf.getvalue()

with st.sidebar:
    st.header('⚙️ Hesaplama ayarları')
    work_rate=st.number_input('Asgari işçilik oranı (%)',min_value=0.,max_value=100.,value=9.0,step=.25,help='Yapı türüne ait resmî SGK tebliğ oranını kontrol ederek giriniz. %38,75 burada kullanılmaz.')
    discount=st.number_input('Araştırma indirimi (%)',min_value=0.,max_value=100.,value=25.,step=1.,help='SGK araştırmasının niteliğine göre uygulanabilmesi kontrol edilmelidir.')
    premium_rate=st.number_input('SGK toplam prim oranı (%)',min_value=0.,max_value=100.,value=38.75,step=.25)
    st.warning('2026 %38,75 SGK toplam prim oranıdır; asgari işçilik oranı değildir.')
    st.caption('Bu hesap ön değerlendirmedir; SGK’nın nihai tahakkuku değildir.')

st.subheader('1 · Proje bilgileri')
c1,c2,c3=st.columns([2,1,1]);name=c1.text_input('Proje / mükellef adı',placeholder='Örn. 34 Villa Projesi');ada=c2.text_input('Ada No');parsel=c3.text_input('Parsel No')
c1,c2,c3=st.columns(3)
start=c1.date_input('İnşaat başlangıç tarihi',value=date(2024,12,1),format='DD.MM.YYYY')
end=c2.date_input('İnşaat bitiş / hesap tarihi',value=date.today(),format='DD.MM.YYYY')
c3.info('Tarife yılı, başlangıç ve bitiş tarihlerine göre otomatik seçilir.')
if end<start:st.error('Bitiş tarihi başlangıçtan önce olamaz.');st.stop()

st.subheader('2 · Yapı bilgileri')
classes=['I-A','I-B','I-C','I-D','II-A','II-B','II-C','III-A','III-B','III-C','IV-A','IV-B','IV-C','V-A','V-B','V-C','V-D','V-E']
if 'blocks' not in st.session_state:st.session_state.blocks=[]
if 'next_id' not in st.session_state:st.session_state.next_id=1

def add_block():
    i=st.session_state.next_id;st.session_state.next_id+=1;st.session_state.blocks.append(i)

def remove_block(i):st.session_state.blocks.remove(i)

st.caption('Blok yoksa doğrudan tek yapı alanını doldur. Varsa aşağıdaki + butonuyla istediğin kadar blok ekle.')
if not st.session_state.blocks:
    x,y=st.columns(2);area=x.number_input('Toplam inşaat alanı (m²)',min_value=0.0,step=1.,format='%.2f');klass=y.selectbox('Yapı sınıfı',classes,index=classes.index('IV-B'))
    entries=[{'name':'Tek yapı','area':area,'klass':klass,'start':start,'end':end}]
else:
    entries=[]
    for i in st.session_state.blocks:
        with st.container(border=True):
            a,b,c,d=st.columns([2,2,2,.7]); n=a.text_input('Blok adı',value=f'{i}. Blok',key=f'n{i}');ar=b.number_input('Alan (m²)',min_value=0.,format='%.2f',key=f'a{i}');cl=c.selectbox('Yapı sınıfı',classes,index=classes.index('IV-B'),key=f'c{i}');
            if d.button('🗑️',key=f'd{i}',help='Bloğu sil'):remove_block(i);st.rerun()
            x,y=st.columns(2);bs=x.date_input('Başlangıç',value=start,key=f's{i}',format='DD.MM.YYYY');be=y.date_input('Bitiş',value=end,key=f'e{i}',format='DD.MM.YYYY')
            entries.append({'name':n,'area':ar,'klass':cl,'start':bs,'end':be})
if st.button('➕ Blok Ekle',on_click=add_block):st.rerun()

st.subheader('3 · Otomatik birim maliyet tarifesi')
st.caption('Aynı yıl veya takip eden yıl bitişinde başlangıç yılı; 3 yılda aradaki 2 yıl; 4 yılda aradaki 3 yılın ortalaması uygulanır. Tanımlı tarifeler: 2024–2026.')

# Kullanıcının seçtiği mevzuat yöntemi tarife yıllarını belirler.
needed = sorted({(y, e['klass']) for e in entries for y in tariff_years(e)})
missing_rates = [(y, cl) for y, cl in needed if tariff_cost(y, cl) is None]
overrides = {}
if missing_rates:
    st.warning('Tarifede bulunamayan yıl/sınıf kombinasyonları için resmî tutarı kontrol ederek giriniz.')
    for year, cl in missing_rates:
        overrides[(year, cl)] = st.number_input(f'{year} / {cl} — TL/m² (elle doğrulanacak)', min_value=0., step=100., key=f'override_{year}_{cl}')

for year, cl in needed:
    amount = tariff_cost(year, cl)
    if amount is not None:
        st.write(f'✅ **{year} / {cl}: {tr(amount, " TL/m²")}** (resmî tarife)')

def unit_cost(e):
    years = tariff_years(e)
    values = [tariff_cost(y, e['klass']) if tariff_cost(y, e['klass']) is not None else overrides.get((y, e['klass']), 0) for y in years]
    if not all(values):
        return None, years
    return sum(values)/len(values), years

st.subheader('4 · Aylık SGK bildirimi / Excel')
months=['Ocak','Şubat','Mart','Nisan','Mayıs','Haziran','Temmuz','Ağustos','Eylül','Ekim','Kasım','Aralık']
monthly_base=[[y,m,0.,0.] for y in range(start.year,end.year+1) for mi,m in enumerate(months,1) if date(y,mi,1)<=end and date(y,mi,calendar.monthrange(y,mi)[1])>=start]
if 'monthly_data' not in st.session_state:st.session_state.monthly_data=monthly_base
if 'monthly_sig' not in st.session_state:st.session_state.monthly_sig=None
sig=(start,end)
if st.session_state.monthly_sig!=sig:
    prev={(r[0],r[1]):r for r in st.session_state.monthly_data}
    st.session_state.monthly_data=[prev.get((r[0],r[1]),r) for r in monthly_base]
    st.session_state.monthly_sig=sig
x,y=st.columns(2)
with x:
    st.info('Aşağıdaki hesap özetiyle aynı sayfada, alt kısmında aylık SGK satırları olan Excel sonuç bölümünden indirilebilir.')
with y:
    upload=st.file_uploader('Doldurulmuş Excel’i yükle',type=['xlsx'],label_visibility='collapsed')
if upload is not None:
    filebytes=upload.getvalue();finger=(upload.name,len(filebytes),__import__('hashlib').md5(filebytes).hexdigest())
    if st.session_state.get('last_upload')!=finger:
        try:
            loaded=read_xlsx(filebytes)
            known={(r[0],r[1]):r for r in loaded}
            st.session_state.monthly_data=[known.get((r[0],r[1]),r) for r in monthly_base]
            st.session_state.last_upload=finger
            st.success(f'{len(loaded)} aylık kayıt okundu.')
        except Exception as exc:st.error(f'Excel okunamadı: {exc}')
frame=pd.DataFrame(st.session_state.monthly_data,columns=['Yıl','Ay','Prime Esas Kazanç (TL)','Ödenen SGK Primi (TL)'])
edited=st.data_editor(frame,hide_index=True,use_container_width=True,disabled=['Yıl','Ay'],column_config={'Prime Esas Kazanç (TL)':st.column_config.NumberColumn(min_value=0,format='%.2f'),'Ödenen SGK Primi (TL)':st.column_config.NumberColumn(min_value=0,format='%.2f')},key='sgk_editor')
monthly=[[int(r['Yıl']),str(r['Ay']),money(r['Prime Esas Kazanç (TL)']),money(r['Ödenen SGK Primi (TL)'])] for _,r in edited.iterrows()]
st.session_state.monthly_data=monthly

st.subheader('5 · ANLIK HESAPLAMA SONUCU')
details=[];warnings=[]
for e in entries:
    c,yrs=unit_cost(e)
    if e['end']<e['start']:warnings.append(f'{e["name"]}: tarih aralığı hatalı')
    if c is None:warnings.append(f'{e["name"]}: {e["klass"]} için {", ".join(map(str,yrs))} yılı/yılları birim maliyeti eksik')
    details.append({'name':e['name'],'area':e['area'],'klass':e['klass'],'cost':c or 0,'total':e['area']*(c or 0),'years':yrs})
for w in warnings:st.warning(w)
all_area=sum(r['area'] for r in details);total=sum(r['total'] for r in details)
required=total*(work_rate/100)*(1-discount/100)
base=sum(r[2] for r in monthly);paid=sum(r[3] for r in monthly)
missing=max(0,required-base);surplus=max(0,base-required)
estimated=missing*premium_rate/100
p,q,r,s=st.columns(4)
p.metric('Toplam m²',tr(all_area,' m²'));q.metric('İnşaat maliyeti',tr(total));r.metric('Asgari işçilik matrahı',tr(required));s.metric('Bildirilen PEK',tr(base))
p,q,r=st.columns(3)
p.metric('EKSİK İŞÇİLİK',tr(missing));q.metric('FAZLA BİLDİRİM',tr(surplus));r.metric('TAHMİNİ EK PRİM',tr(estimated))
st.caption('Ödenen SGK primi: '+tr(paid)+' | Asgari işçilik kıyası ödenen primle değil, prime esas kazançla yapılır.')
if warnings:st.error('Birim maliyetler tamamlanmadığı için hesap ve PDF sonucu kesinleştirilmemiştir. Önce eksikleri doldurunuz.')
st.dataframe(pd.DataFrame(details).rename(columns={'name':'Yapı','area':'m²','klass':'Sınıf','cost':'Birim TL/m²','total':'Toplam TL','years':'Tarife yılları'}),hide_index=True,use_container_width=True)
project={'ad':name or 'Belirtilmedi','ada':ada or '-','parsel':parsel or '-','dates':f'{start:%d.%m.%Y} - {end:%d.%m.%Y}','discount':discount}
try:
    xlsx=build_xlsx(project,details,monthly,total,required,base,missing,estimated,work_rate,premium_rate)
    st.download_button('📊 EXCEL İNDİR (ÖZET + AYLIK SGK TAKİP)',data=xlsx,file_name=f'DEVRAN_SGK_TAKIP_{ada or "ADA"}_{parsel or "PARSEL"}.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',disabled=bool(warnings) or all_area<=0,use_container_width=True)
except Exception as e:st.error(f'Excel oluşturulamadı: {e}')
try:
    pdf=pdf_report(project,details,monthly,total,required,base,missing,estimated,work_rate,premium_rate)
    st.download_button('📄 PDF İNDİR (ÖZET SONUÇ)',data=pdf,file_name=f'DEVRAN_ASGARI_ISCILIK_{ada or "ADA"}_{parsel or "PARSEL"}.pdf',mime='application/pdf',disabled=bool(warnings) or all_area<=0,type='primary',use_container_width=True)
except Exception as e:st.error(f'PDF oluşturulamadı: {e}')
st.caption('Not: Tarife seçimi istenen kurala göre yapılmaktadır; özel durumlar ve SGK mevzuatı ayrıca kontrol edilmelidir. Kısmi ilişiksizlik dönem ayrımı sonraki sürüme eklenecektir.')
